import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

wb = openpyxl.Workbook()

# =============================================
# 共通スタイル定義
# =============================================
THIN = Side(style="thin", color="AAAAAA")
MEDIUM = Side(style="medium", color="555555")
BORDER_ALL_THIN = Border(top=THIN, left=THIN, right=THIN, bottom=THIN)
BORDER_ALL_MEDIUM = Border(top=MEDIUM, left=MEDIUM, right=MEDIUM, bottom=MEDIUM)

COLOR_HEADER_DARK  = "1F4E79"
COLOR_SECTION_A    = "2E75B6"
COLOR_SECTION_B    = "ED7D31"
COLOR_SECTION_C    = "70AD47"
COLOR_SECTION_D    = "4472C4"
COLOR_SECTION_E    = "C00000"
COLOR_SECTION_F    = "7030A0"
COLOR_SECTION_G    = "404040"
COLOR_LABEL_BG     = "D9E2F3"
COLOR_INPUT_BG     = "FAFAFA"
COLOR_REQUIRED     = "FFF2CC"
COLOR_WHITE        = "FFFFFF"
COLOR_TITLE_FG     = "FFFFFF"

def make_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def set_cell(ws, row, col, value="", font_bold=False, font_color="000000",
             font_size=10, bg_color=None, h_align="left", v_align="center",
             wrap=True, border=BORDER_ALL_THIN):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=font_bold, color=font_color, size=font_size, name="Yu Gothic")
    cell.alignment = Alignment(horizontal=h_align, vertical=v_align, wrap_text=wrap)
    if bg_color:
        cell.fill = make_fill(bg_color)
    if border:
        cell.border = border
    return cell

def section_header(ws, row, col_start, col_end, label, color):
    ws.merge_cells(start_row=row, start_column=col_start,
                   end_row=row, end_column=col_end)
    cell = ws.cell(row=row, column=col_start, value=f"  {label}")
    cell.font = Font(bold=True, color=COLOR_TITLE_FG, size=11, name="Yu Gothic")
    cell.fill = make_fill(color)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    cell.border = BORDER_ALL_MEDIUM
    ws.row_dimensions[row].height = 22

def input_row(ws, row, label, height=45, required=False, note=""):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    label_text = f"  {'★ ' if required else ''}{label}"
    set_cell(ws, row, 1, label_text, font_bold=True, font_size=10,
             bg_color=COLOR_LABEL_BG, h_align="left")
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
    input_bg = COLOR_REQUIRED if required else COLOR_INPUT_BG
    set_cell(ws, row, 3, "", bg_color=input_bg, h_align="left", wrap=True)
    set_cell(ws, row, 7, note, font_size=9, bg_color=COLOR_WHITE,
             h_align="left", wrap=True)
    ws.row_dimensions[row].height = height

def rating_row(ws, row, label, required=False):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    label_text = f"  {'★ ' if required else ''}{label}"
    set_cell(ws, row, 1, label_text, font_bold=True, font_size=10,
             bg_color=COLOR_LABEL_BG)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
    input_bg = COLOR_REQUIRED if required else COLOR_INPUT_BG
    set_cell(ws, row, 3, "", bg_color=input_bg, h_align="left", wrap=True)

    cell_e = ws.cell(row=row, column=5)
    cell_e.value = "▼ 評価を選択"
    cell_e.font = Font(size=9, color="888888", name="Yu Gothic", italic=True)
    cell_e.fill = make_fill("EAF0FB")
    cell_e.alignment = Alignment(horizontal="center", vertical="center")
    cell_e.border = BORDER_ALL_THIN

    dv = DataValidation(
        type="list",
        formula1='"予定通り,一部未達,未達,課題あり,該当なし"',
        allow_blank=True,
        showDropDown=False,
        error="リストから選択してください",
        errorTitle="入力エラー",
        prompt="評価を選択してください",
        promptTitle="評価"
    )
    ws.add_data_validation(dv)
    dv.add(cell_e)

    ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=7)
    set_cell(ws, row, 6, "", bg_color=COLOR_INPUT_BG, wrap=True)
    ws.row_dimensions[row].height = 50


# =============================================
# シート1：業務報告票
# =============================================
ws1 = wb.active
ws1.title = "業務報告票"

col_widths = {1: 4, 2: 22, 3: 28, 4: 28, 5: 16, 6: 16, 7: 22}
for col, width in col_widths.items():
    ws1.column_dimensions[get_column_letter(col)].width = width

# タイトル
ws1.merge_cells("A1:G1")
set_cell(ws1, 1, 1, "業　務　報　告　票",
         font_bold=True, font_color=COLOR_TITLE_FG, font_size=16,
         bg_color=COLOR_HEADER_DARK, h_align="center", border=BORDER_ALL_MEDIUM)
ws1.row_dimensions[1].height = 36

ws1.merge_cells("A2:G2")
set_cell(ws1, 2, 1,
         "  ★ は必須入力項目です　／　E列「評価」は前回取り組みの結果評価（ドロップダウン）",
         font_size=9, font_color="555555", bg_color="EFF3FB",
         h_align="left", border=BORDER_ALL_THIN)
ws1.row_dimensions[2].height = 16

r = 3

# 0. 基本情報
section_header(ws1, r, 1, 7, "0．基本情報", COLOR_SECTION_A); r += 1

for col, txt in [(1, "項目"), (3, "内容"), (5, "項目"), (6, "内容")]:
    cell = ws1.cell(row=r, column=col, value=txt)
    cell.font = Font(bold=True, size=9, color="555555", name="Yu Gothic")
    cell.fill = make_fill("CFD5EA")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER_ALL_THIN

ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
ws1.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
ws1.merge_cells(start_row=r, start_column=5, end_row=r, end_column=5)
ws1.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
ws1.row_dimensions[r].height = 16; r += 1

basic_pairs = [
    ("氏名", "所属部署"),
    ("報告対象期間", "報告日"),
    ("現場名 / プロジェクト名", "役割"),
]
for left_label, right_label in basic_pairs:
    ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws1.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    ws1.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    set_cell(ws1, r, 1, f"  ★ {left_label}", font_bold=True, font_size=10, bg_color=COLOR_LABEL_BG)
    set_cell(ws1, r, 3, "", bg_color=COLOR_REQUIRED)
    set_cell(ws1, r, 5, f"  ★ {right_label}", font_bold=True, font_size=10, bg_color=COLOR_LABEL_BG)
    set_cell(ws1, r, 6, "", bg_color=COLOR_REQUIRED)
    ws1.row_dimensions[r].height = 22; r += 1

# 1. 携わっている案件
section_header(ws1, r, 1, 7, "1．携わっている案件", COLOR_SECTION_A); r += 1
input_row(ws1, r, "案件名（複数ある場合は列挙）", height=50, required=True, note="複数案件は箇条書きで記入"); r += 1
input_row(ws1, r, "各案件の進捗状況", height=50, note="進行中／完了／遅延 等を記入"); r += 1

# 2. 前回報告からの結果
section_header(ws1, r, 1, 7, "2．前回報告からの結果（PDCA：Check ／ KPT：Problem）", COLOR_SECTION_B); r += 1

for col, txt, span_end in [(1,"中項目",2),(3,"内容",4),(5,"結果評価",5),(6,"補足・備考",7)]:
    ws1.merge_cells(start_row=r, start_column=col, end_row=r, end_column=span_end)
    cell = ws1.cell(row=r, column=col, value=txt)
    cell.font = Font(bold=True, size=9, color="555555", name="Yu Gothic")
    cell.fill = make_fill("F4C7A8")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER_ALL_THIN
ws1.row_dimensions[r].height = 16; r += 1

rating_row(ws1, r, "前回設定した取り組み内容", required=True); r += 1
rating_row(ws1, r, "実施結果（何をどこまで実施したか）", required=True); r += 1
rating_row(ws1, r, "成果・達成できたこと"); r += 1
rating_row(ws1, r, "未達・残課題（未達の場合その理由）"); r += 1
input_row(ws1, r, "次回報告への反映事項", height=45, note="未達内容を次回計画に転記する"); r += 1

# 3. 現在の作業内容・状況
section_header(ws1, r, 1, 7, "3．現在の作業内容・状況（PDCA：Do）", COLOR_SECTION_C); r += 1
input_row(ws1, r, "実施している主な作業内容", height=60, required=True, note="箇条書き推奨"); r += 1
input_row(ws1, r, "現在の進捗・状況", height=45, required=True, note="○%完了、課題対応中 等"); r += 1
input_row(ws1, r, "成果物・変更点", height=45, note="ドキュメント・設定変更など"); r += 1

# 4. 現場での気づき
section_header(ws1, r, 1, 7, "4．現場での気づき・ナレッジ共有（KPT：Keep）", COLOR_SECTION_D); r += 1
input_row(ws1, r, "気づき・学んだこと", height=60, note="再発防止・効率化のヒントなど"); r += 1
input_row(ws1, r, "他メンバーへ共有したいナレッジ", height=60, note="横展開できるノウハウを記入"); r += 1

# 5. 困りごと・懸念事項
section_header(ws1, r, 1, 7, "5．現場での困りごと・懸念事項（KPT：Problem）", COLOR_SECTION_E); r += 1
input_row(ws1, r, "困りごと（工数・体制・ツール等）", height=60, note="具体的に記入（5W1H）"); r += 1
input_row(ws1, r, "懸念事項（納期・品質・リスク等）", height=60, note="影響度・緊急度も添えると◎"); r += 1
input_row(ws1, r, "会社・上長に支援してほしいこと", height=45); r += 1

# 6. 次回に向けた取り組み
section_header(ws1, r, 1, 7, "6．次回に向けた取り組み（PDCA：Plan・Act ／ KPT：Try）", COLOR_SECTION_F); r += 1
input_row(ws1, r, "次回の取り組み目標（Plan）", height=50, required=True, note="具体的・測定可能な目標"); r += 1
input_row(ws1, r, "実施予定の作業・期限（Do予定）", height=60, required=True, note="担当・期限も記入"); r += 1
input_row(ws1, r, "改善策・工夫点（Act）", height=50, note="前回の未達を踏まえた改善"); r += 1
input_row(ws1, r, "リスク・懸念（Plan時点）", height=45); r += 1

# 7. 上長コメント
section_header(ws1, r, 1, 7, "7．上長コメント・フィードバック", COLOR_SECTION_G); r += 1
input_row(ws1, r, "コメント・指示事項", height=70, note="報告受領後に上長が記入"); r += 1

ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
set_cell(ws1, r, 1, "  確認日：　　　　　　　　　　　確認者氏名：",
         font_size=10, bg_color="EEEEEE", h_align="left")
ws1.row_dimensions[r].height = 22

ws1.page_setup.orientation = "portrait"
ws1.page_setup.paperSize = 9
ws1.page_margins.left = 0.5
ws1.page_margins.right = 0.5
ws1.page_margins.top = 0.75
ws1.page_margins.bottom = 0.75
ws1.print_title_rows = "$1:$2"


# =============================================
# シート2：次回アクション管理
# =============================================
ws2 = wb.create_sheet("次回アクション管理")

col_widths2 = {1:6, 2:20, 3:30, 4:14, 5:14, 6:20, 7:30, 8:20}
for col, width in col_widths2.items():
    ws2.column_dimensions[get_column_letter(col)].width = width

ws2.merge_cells("A1:H1")
set_cell(ws2, 1, 1, "次回アクション管理シート",
         font_bold=True, font_color=COLOR_TITLE_FG, font_size=14,
         bg_color=COLOR_HEADER_DARK, h_align="center", border=BORDER_ALL_MEDIUM)
ws2.row_dimensions[1].height = 32

ws2.merge_cells("A2:H2")
set_cell(ws2, 2, 1,
         "  このシートは「業務報告票 ▶ 6．次回に向けた取り組み」の内容を転記し、次回報告時にフォローします。",
         font_size=9, font_color="555555", bg_color="EFF3FB",
         h_align="left", border=BORDER_ALL_THIN)
ws2.row_dimensions[2].height = 16

headers2 = ["No", "氏名", "取り組み内容（Plan）", "予定完了日",
            "実際完了日", "結果評価", "実施結果・コメント", "次回への引き継ぎ"]
header_colors2 = ["1F4E79","1F4E79","2E75B6","2E75B6",
                  "ED7D31","ED7D31","70AD47","7030A0"]
for col, (hdr, color) in enumerate(zip(headers2, header_colors2), 1):
    cell = ws2.cell(row=3, column=col, value=hdr)
    cell.font = Font(bold=True, color=COLOR_TITLE_FG, size=10, name="Yu Gothic")
    cell.fill = make_fill(color)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER_ALL_MEDIUM
ws2.row_dimensions[3].height = 30

dv2 = DataValidation(
    type="list",
    formula1='"予定通り,一部未達,未達,課題あり,該当なし"',
    allow_blank=True,
    showDropDown=False
)
ws2.add_data_validation(dv2)

for i in range(1, 16):
    row_idx = 3 + i
    row_colors = [
        "F2F2F2", COLOR_INPUT_BG, COLOR_INPUT_BG, "FFF2CC",
        "FFF2CC", "EAF0FB", COLOR_INPUT_BG, COLOR_INPUT_BG
    ]
    for col in range(1, 9):
        cell = ws2.cell(row=row_idx, column=col,
                        value=str(i) if col == 1 else "")
        cell.font = Font(size=10, name="Yu Gothic")
        cell.fill = make_fill(row_colors[col - 1])
        cell.alignment = Alignment(horizontal="left" if col != 1 else "center",
                                   vertical="center", wrap_text=True)
        cell.border = BORDER_ALL_THIN
    dv2.add(ws2.cell(row=row_idx, column=6))
    ws2.row_dimensions[row_idx].height = 40

ws2.freeze_panes = "A4"
ws2.page_setup.orientation = "landscape"
ws2.page_setup.paperSize = 9
ws2.page_margins.left = 0.4
ws2.page_margins.right = 0.4

# =============================================
# 保存
# =============================================
output_path = "業務報告書_テンプレート.xlsx"
wb.save(output_path)
print(f"Excelファイルを作成しました：{output_path}")
